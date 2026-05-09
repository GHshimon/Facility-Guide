from openpyxl import Workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = "D:/Private/shimon/Project/相談員アプリ/施設データ.xlsx"

HEADER_FILL   = PatternFill("solid", start_color="1F4E79")
SUB_FILL      = PatternFill("solid", start_color="2E75B6")
CARE_FILL     = PatternFill("solid", start_color="375623")
ECON_FILL     = PatternFill("solid", start_color="7B3F00")
INFO_FILL     = PatternFill("solid", start_color="4B0082")
ALT_ROW_FILL  = PatternFill("solid", start_color="F2F7FF")
REQ_FILL      = PatternFill("solid", start_color="FFF2CC")

WHITE_BOLD = Font(name="Arial", bold=True, color="FFFFFF", size=10)
WHITE_FONT = Font(name="Arial", color="FFFFFF", size=9)
BODY_FONT  = Font(name="Arial", size=9)
NOTE_FONT  = Font(name="Arial", size=8, italic=True, color="595959")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ─────────────────────────────────────────
# Sheet 1: 施設一覧
# ─────────────────────────────────────────
ws = wb.active
ws.title = "施設一覧"
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A4"

# ── Section header rows ──
# Row 1: Title
ws.merge_cells("A1:AH1")
ws["A1"] = "施設データベース（鹿児島県）"
ws["A1"].font = Font(name="Arial", bold=True, size=14, color="1F4E79")
ws["A1"].alignment = LEFT
ws["A1"].fill = PatternFill("solid", start_color="DEEAF1")
ws.row_dimensions[1].height = 28

# Row 2: Category headers (merged)
category_defs = [
    ("A2", "C2", "基本情報",     "1F4E79"),
    ("D2", "H2", "受入可能要介護度",  "2E75B6"),
    ("I2", "J2", "部屋",         "2E75B6"),
    ("K2", "Q2", "医療ケア対応",  "375623"),
    ("R2", "V2", "経済条件",      "7B3F00"),
    ("W2", "AB2","連絡先",        "404040"),
    ("AC2","AH2","相談員生情報",  "4B0082"),
]
for start, end, label, color in category_defs:
    ws.merge_cells(f"{start}:{end}")
    c = ws[start]
    c.value = label
    c.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
    c.fill = PatternFill("solid", start_color=color)
    c.alignment = CENTER
ws.row_dimensions[2].height = 20

# Row 3: Column headers
col_headers = [
    # 基本情報
    "ID", "施設名", "施設種別",
    # 受入要介護度
    "要介護1", "要介護2", "要介護3", "要介護4", "要介護5",
    # 部屋
    "個室", "多床室",
    # 医療ケア
    "インスリン", "胃瘻", "経管栄養", "喀痰吸引",
    "バルーン", "褥瘡処置", "透析",
    # 経済条件
    "生保受入", "第1段階", "第2段階", "第3段階", "第4段階",
    # 連絡先
    "エリア", "住所", "電話", "FAX", "担当者名", "担当者メモ",
    # 生情報
    "相談員柔軟性", "空床傾向", "判定会議頻度",
    "施設の強み", "月額費用目安", "備考（生情報）",
]

col_colors = (
    ["1F4E79"]*3 +
    ["2E75B6"]*5 +
    ["2E75B6"]*2 +
    ["375623"]*7 +
    ["7B3F00"]*5 +
    ["404040"]*6 +
    ["4B0082"]*6
)

for i, (header, color) in enumerate(zip(col_headers, col_colors), start=1):
    c = ws.cell(row=3, column=i, value=header)
    c.font = WHITE_BOLD
    c.fill = PatternFill("solid", start_color=color)
    c.alignment = CENTER
    c.border = BORDER
ws.row_dimensions[3].height = 32

# ── Column widths ──
col_widths = [
    5, 22, 14,          # 基本
    7, 7, 7, 7, 7,      # 要介護度
    7, 7,               # 部屋
    10, 8, 10, 10, 9, 10, 8,  # 医療ケア
    9, 8, 8, 8, 8,      # 経済
    10, 28, 16, 16, 12, 20,   # 連絡先
    12, 12, 14, 20, 20, 36,   # 生情報
]
for i, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ── Data validation ──
OX_VALS  = '"○,×,要相談"'
FLEX_VALS = '"高,中,低"'
VAC_VALS  = '"あり,やや少ない,なし,不明"'
BOOL_VALS = '"○,×"'
TYPE_VALS = '"特別養護老人ホーム,介護老人保健施設,介護医療院,サービス付き高齢者向け住宅,特定施設入居者生活介護,グループホーム,その他"'
AREA_VALS = '"鹿児島市,霧島市,姶良市,北薩,南薩,大隅,離島"'

def dv(formula, col_letter, start=4, end=200):
    v = DataValidation(type="list", formula1=formula, allow_blank=True,
                       sqref=f"{col_letter}{start}:{col_letter}{end}")
    ws.add_data_validation(v)

dv(TYPE_VALS, "C")
for col in ["D","E","F","G","H","I","J"]:
    dv(BOOL_VALS, col)
for col in ["K","L","M","N","O","P","Q"]:
    dv(OX_VALS, col)
for col in ["R","S","T","U","V"]:
    dv(BOOL_VALS, col)
dv(AREA_VALS,  "W")
dv(FLEX_VALS,  "AC")
dv(VAC_VALS,   "AD")

# ── Sample data ──
sample = [
    [1, "さくら苑", "特別養護老人ホーム",
     "×","×","○","○","○",
     "○","○",
     "○","○","○","○","○","要相談","×",
     "○","○","○","○","○",
     "鹿児島市","鹿児島市吉野町1234","099-xxx-0001","099-xxx-0002","田中 花子","明るく丁寧",
     "高","あり","週1回","認知症対応・看取り対応","第1段階:約3万円","BPSDのある方も受け入れ実績あり。相談員の対応が丁寧。"],
    [2, "霧島老人ホーム", "特別養護老人ホーム",
     "×","×","○","○","○",
     "○","×",
     "○","要相談","要相談","○","○","○","×",
     "○","○","○","○","○",
     "霧島市","霧島市国分中央5678","0995-xxx-0003","0995-xxx-0004","山田 一郎","やや慎重",
     "中","やや少ない","月2回","看取り対応","第1段階:約3万円","判定まで時間がかかることがある。早めの打診推奨。"],
    [3, "姶良リハビリ老健", "介護老人保健施設",
     "○","○","○","○","○",
     "○","○",
     "○","○","○","○","○","○","×",
     "×","×","○","○","○",
     "姶良市","姶良市加治木町9012","0995-xxx-0005","0995-xxx-0006","佐藤 健","柔軟",
     "高","あり","週2回","リハビリ・在宅復帰・医療ケア","第2段階:約4万円","医療依存度高い方も積極受け入れ。在宅復帰が前提。"],
    [4, "南薩介護医療院", "介護医療院",
     "×","×","○","○","○",
     "○","○",
     "○","○","○","○","○","○","要相談",
     "○","○","○","○","○",
     "南薩","南さつま市加世田1357","0993-xxx-0007","0993-xxx-0008","鈴木 明","非常に柔軟",
     "高","あり","週1回","医療ケア全般・難病・看取り","第1段階:約3.5万円","医師常駐。医療依存度の高い方の長期受け入れが得意。"],
    [5, "あおぞら サ高住", "サービス付き高齢者向け住宅",
     "○","○","○","○","×",
     "○","×",
     "○","×","×","要相談","○","要相談","×",
     "×","×","×","○","○",
     "鹿児島市","鹿児島市与次郎2468","099-xxx-0009","099-xxx-0010","木村 智","普通",
     "中","あり","随時","自立度高め・プライバシー重視","月額15〜20万円","賃貸契約のため生保不可。認知症軽度まで対応。"],
    [6, "ひまわりの里", "特別養護老人ホーム",
     "×","×","○","○","○",
     "×","○",
     "○","○","○","要相談","○","要相談","×",
     "○","○","○","○","○",
     "大隅","鹿屋市寿町3579","0994-xxx-0011","0994-xxx-0012","高橋 良子","柔軟",
     "高","あり","月2回","看取り・生保対応実績豊富","第1段階:約2.5万円","地域密着型。エリア外の方でも相談に乗ってくれる。"],
    [7, "北薩ケアホーム", "特定施設入居者生活介護",
     "×","○","○","○","○",
     "○","×",
     "○","要相談","×","×","○","要相談","×",
     "×","×","×","○","○",
     "北薩","薩摩川内市隈之城町2468","0996-xxx-0013","0996-xxx-0014","伊藤 誠","やや慎重",
     "低","やや少ない","月1回","個室・家族サービス充実","月額16〜22万円","経済的余裕のある方向け。入居審査がやや厳しい。"],
    [8, "鹿児島市内 老健センター", "介護老人保健施設",
     "○","○","○","○","○",
     "○","○",
     "○","○","○","○","○","○","×",
     "○","○","○","○","○",
     "鹿児島市","鹿児島市和田1-23-4","099-xxx-0015","099-xxx-0016","渡辺 香","丁寧",
     "高","やや少ない","週1回","リハビリ・短期集中リハ・医療ケア","第1段階:約3万円","在宅復帰率が高くリハビリ目的での入所に向いている。"],
]

OX_GREEN  = PatternFill("solid", start_color="E8F5E9")
OX_YELLOW = PatternFill("solid", start_color="FFF9C4")
OX_RED    = PatternFill("solid", start_color="FFEBEE")

def cell_fill_for_ox(val):
    if val == "○": return OX_GREEN
    if val == "要相談": return OX_YELLOW
    if val == "×": return OX_RED
    return None

for r_idx, row in enumerate(sample, start=4):
    fill = ALT_ROW_FILL if r_idx % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
    for c_idx, val in enumerate(row, start=1):
        c = ws.cell(row=r_idx, column=c_idx, value=val)
        c.font = BODY_FONT
        c.border = BORDER
        c.alignment = CENTER if c_idx not in [2, 24, 25, 26, 27, 28, 29, 32, 33, 34, 35] else LEFT
        ox_fill = cell_fill_for_ox(val) if isinstance(val, str) else None
        c.fill = ox_fill if ox_fill else fill
    ws.row_dimensions[r_idx].height = 18

# ── Instruction row (row after data) ──
ws.merge_cells("A203:AH203")
ws["A203"] = "※ 医療ケア欄は「○」「×」「要相談」の3種類で入力してください。要介護度・部屋・生保欄は「○」「×」。"
ws["A203"].font = NOTE_FONT
ws["A203"].fill = PatternFill("solid", start_color="F0F4FF")
ws["A203"].alignment = LEFT

# ─────────────────────────────────────────
# Sheet 2: 入力ガイド
# ─────────────────────────────────────────
wg = wb.create_sheet("入力ガイド")
wg.sheet_view.showGridLines = False
wg.column_dimensions["A"].width = 22
wg.column_dimensions["B"].width = 50
wg.column_dimensions["C"].width = 30

wg.merge_cells("A1:C1")
wg["A1"] = "入力ガイド"
wg["A1"].font = Font(name="Arial", bold=True, size=13, color="1F4E79")
wg["A1"].fill = PatternFill("solid", start_color="DEEAF1")
wg["A1"].alignment = LEFT
wg.row_dimensions[1].height = 26

guide = [
    ("項目", "説明", "入力例"),
    ("施設種別", "施設の種類をドロップダウンから選択", "特別養護老人ホーム"),
    ("要介護1〜5", "その要介護度の受け入れが可能かを ○/× で入力", "○"),
    ("個室/多床室", "その部屋タイプがあるかを ○/× で入力", "○"),
    ("医療ケア各項目", "○：対応可　×：対応不可　要相談：条件次第", "要相談"),
    ("生保受入", "生活保護受給者の受け入れ可否を ○/× で入力", "○"),
    ("第1〜4段階", "負担限度額認定の各段階に対応しているか ○/× で入力", "○"),
    ("エリア", "ドロップダウンから選択\n（鹿児島市/霧島市/姶良市/北薩/南薩/大隅/離島）", "鹿児島市"),
    ("相談員柔軟性", "相談員の対応の柔軟さを 高/中/低 で評価", "高"),
    ("空床傾向", "あり/やや少ない/なし/不明 から選択", "あり"),
    ("判定会議頻度", "判定会議の頻度を自由入力", "週1回"),
    ("施設の強み", "得意なケア・特徴をカンマ区切りで入力", "認知症対応,看取り"),
    ("月額費用目安", "段階別の目安金額", "第1段階:約3万円"),
    ("備考（生情報）", "相談員だけが知る非公式情報をメモ", "担当の田中さんは柔軟に対応してくれる"),
]

for r_idx, row in enumerate(guide, start=2):
    for c_idx, val in enumerate(row, start=1):
        c = wg.cell(row=r_idx, column=c_idx, value=val)
        if r_idx == 2:
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
            c.fill = PatternFill("solid", start_color="1F4E79")
        else:
            c.font = Font(name="Arial", size=9)
            c.fill = ALT_ROW_FILL if r_idx % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = BORDER
    wg.row_dimensions[r_idx].height = 28

# ─────────────────────────────────────────
# Save
# ─────────────────────────────────────────
wb.save(OUTPUT)
print("Done: " + OUTPUT)
