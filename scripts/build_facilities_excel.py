# -*- coding: utf-8 -*-
"""
鹿児島県施設データ Excel 生成スクリプト
  - 厚労省 介護サービス情報公表オープンデータ（2025年12月末）
  - 鹿児島県 介護保険指定事業所一覧（R7.10.01）
  - Web検索による月額費用目安（施設種別ごとの相場、2025年調査）
を統合して 施設データ.xlsx の行4以降を上書きする
"""

import csv
import os
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

import openpyxl
from openpyxl import load_workbook

# ── パス定義 ──────────────────────────────────────
BASE = "D:/Private/shimon/Project/相談員アプリ"
MHLW_DIR = f"{BASE}/data/公的データ/厚生労働省_介護サービス情報公表オープンデータ_2025年12月末"
PREF_DIR = f"{BASE}/data/公的データ/鹿児島県_介護保険指定事業所一覧_R07.10.01"
PREF_SISETU = f"{PREF_DIR}/R07.10.01_令和7年10月1日現在（居宅サービス・施設サービス・居宅介護支援）（EXCEL：379KB）.xlsx"
PREF_CHIIKI = f"{PREF_DIR}/R07.10.01_令和7年10月1日現在（地域密着型サービス）（EXCEL：156KB）.xlsx"
OUT_XLSX    = f"{BASE}/施設データ.xlsx"

# ── CSVファイル → 施設種別ラベル ──────────────────
CSV_FILES = {
    "510_介護老人福祉施設.csv":                          "特別養護老人ホーム",
    "520_介護老人保健施設.csv":                          "介護老人保健施設",
    "540_地域密着型介護老人福祉施設入所者生活介護.csv":  "地域密着型特養（29人以下）",
    "550_介護医療院.csv":                                "介護医療院",
    "331_特定施設入居者生活介護_有料老人ホーム.csv":     "有料老人ホーム（特定施設）",
    "332_特定施設入居者生活介護_軽費老人ホーム.csv":     "軽費老人ホーム（特定施設）",
    "334_特定施設入居者生活介護_サービス付き高齢者向け住宅.csv": "サービス付き高齢者向け住宅（特定施設）",
}

# ── Web検索で取得した施設種別ごとの月額費用目安 ───
# 出典: LIFULL介護・ケアスル介護 2025年調査（鹿児島県相場）
COST_BY_TYPE = {
    "特別養護老人ホーム":           "多床室7〜9万円、ユニット型個室10〜15万円（要介護度・負担段階により異なる）",
    "介護老人保健施設":             "約8〜15万円（鹿児島県実績5.5〜12.9万円、要介護度・居室により異なる）",
    "地域密着型特養（29人以下）":   "多床室7〜9万円、ユニット型個室10〜15万円（特養に準ずる）",
    "介護医療院":                   "約10〜17万円（医療ケア含む、要介護度・居室により異なる）",
    "有料老人ホーム（特定施設）":   "約10〜25万円（施設・サービス内容により大きく異なる）",
    "軽費老人ホーム（特定施設）":   "約5〜10万円（所得に応じた軽減あり）",
    "サービス付き高齢者向け住宅（特定施設）": "約12〜15万円（鹿児島県平均入居一時金10.2万円・月額12.9万円）",
}
COST_SOURCE = "※月額費用目安はAI Web検索による種別相場（LIFULL介護・ケアスル介護 2025年調査）"

# ── ① 鹿児島県 Excel から 事業所番号→{住所, 定員} を取得 ─
pref_map = {}  # key: 事業所番号(str) → {addr, capacity}

def load_pref_sheet(wb, sheet_name, sisetu_label):
    if sheet_name not in wb.sheetnames:
        return
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=3, values_only=True):
        jigyo_num = str(row[2]).strip().replace(".0", "") if row[2] else ""
        address   = str(row[6]).strip() if row[6] else ""
        capacity_val = row[10]
        capacity = None
        if capacity_val is not None:
            try:
                capacity = int(float(str(capacity_val)))
            except (ValueError, TypeError):
                pass
        if jigyo_num:
            pref_map[jigyo_num] = {"addr": address, "capacity": capacity, "label": sisetu_label}

wb_sisetu = load_workbook(PREF_SISETU, data_only=True)
load_pref_sheet(wb_sisetu, "福祉施設",  "特別養護老人ホーム")
load_pref_sheet(wb_sisetu, "保健施設",  "介護老人保健施設")
load_pref_sheet(wb_sisetu, "介護医療院","介護医療院")

wb_chiiki = load_workbook(PREF_CHIIKI, data_only=True)
load_pref_sheet(wb_chiiki, "地域特養", "地域密着型特養（29人以下）")

print(f"鹿児島県Excelから読み込み: {len(pref_map)}件")

# ── ② 厚労省 CSV を読み込み ───────────────────────
facilities = []

for fname, sisetu_label in CSV_FILES.items():
    fpath = os.path.join(MHLW_DIR, fname)
    with open(fpath, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get("都道府県名", "") != "鹿児島県":
                continue

            jigyo_num  = str(row.get("事業所番号", "")).strip()
            name       = row.get("事業所名", "").strip()
            if not name:
                continue
            city       = row.get("市区町村名", "").strip()
            addr_raw   = row.get("住所", "").strip()
            bldg       = row.get("方書（ビル名等）", "").strip()
            tel        = row.get("電話番号", "").strip()
            fax        = row.get("FAX番号", "").strip()
            corp_name  = row.get("法人の名称", "").strip()
            url        = row.get("URL", "").strip()
            capacity_csv = row.get("定員", "").strip()

            # 住所を完全形に
            if jigyo_num in pref_map and pref_map[jigyo_num]["addr"]:
                full_addr = pref_map[jigyo_num]["addr"]
            else:
                if addr_raw.startswith("鹿児島県") or addr_raw.startswith(city):
                    full_addr = addr_raw if addr_raw.startswith("鹿児島県") else f"鹿児島県{addr_raw}"
                else:
                    full_addr = f"鹿児島県{city}{addr_raw}"
            if bldg:
                full_addr += f" {bldg}"

            # 定員
            if jigyo_num in pref_map and pref_map[jigyo_num]["capacity"]:
                capacity = pref_map[jigyo_num]["capacity"]
            elif capacity_csv and capacity_csv.isdigit() and int(capacity_csv) > 0:
                capacity = int(capacity_csv)
            else:
                capacity = None

            # 検索情報列
            info_parts = [f"事業所番号:{jigyo_num}"]
            if corp_name:
                info_parts.append(f"法人:{corp_name}")
            if capacity:
                info_parts.append(f"定員:{capacity}名")
            if url:
                info_parts.append(f"URL:{url}")
            info_parts.append("出典:厚労省介護サービス情報公表(2025年12月末)・鹿児島県指定事業所一覧(R7.10.01)")
            search_info = " / ".join(info_parts)

            facilities.append({
                "name":        name,
                "type":        sisetu_label,
                "city":        city,
                "address":     full_addr,
                "tel":         tel,
                "fax":         fax,
                "search_info": search_info,
            })

# エリア → 施設名 でソート
facilities.sort(key=lambda x: (x["city"], x["name"]))
print(f"厚労省CSVから取得（鹿児島県絞り込み後）: {len(facilities)}件")

# ── ③ 施設データ.xlsx に書き込む ─────────────────
wb_out = load_workbook(OUT_XLSX)
ws_out = wb_out["施設一覧"]

# 行4以降のマージセルを先に解除（残留マージセルを除去）
for mc in list(ws_out.merged_cells.ranges):
    if mc.min_row >= 4:
        ws_out.unmerge_cells(str(mc))

# 行4以降を削除
if ws_out.max_row >= 4:
    ws_out.delete_rows(4, ws_out.max_row - 3)

# 列ヘッダーを取得（行3）
headers = [ws_out.cell(row=3, column=c).value for c in range(1, ws_out.max_column + 1)]
col_map = {v: i + 1 for i, v in enumerate(headers) if v}
print("列ヘッダー確認:", list(col_map.keys()))

def set_cell(ws, row, col_name, value):
    c = col_map.get(col_name)
    if c:
        ws.cell(row=row, column=c).value = value

# 空の施設名を除外
facilities = [f for f in facilities if f.get("name")]

idx = 0
for f in facilities:
    r = idx + 4
    idx += 1
    set_cell(ws_out, r, "ID",         idx)
    set_cell(ws_out, r, "施設名",     f["name"])
    set_cell(ws_out, r, "施設種別",   f["type"])
    set_cell(ws_out, r, "エリア",     f["city"])
    set_cell(ws_out, r, "住所",       f["address"])
    set_cell(ws_out, r, "電話",       f["tel"])
    set_cell(ws_out, r, "FAX",        f["fax"])
    # 月額費用目安（Web検索で取得した施設種別相場）
    cost_est = COST_BY_TYPE.get(f["type"], "")
    if cost_est:
        set_cell(ws_out, r, "月額費用目安", cost_est)
    # 検索情報（公的データのメタデータ + Web検索出典注記）
    search_info = f["search_info"]
    if cost_est:
        search_info += f" / {COST_SOURCE}"
    set_cell(ws_out, r, "検索情報", search_info)

wb_out.save(OUT_XLSX)
print(f"\n完了: {len(facilities)}件を {OUT_XLSX} に書き込みました。")

# 種別集計
from collections import Counter
wb_check = load_workbook(OUT_XLSX, data_only=True)
ws_check = wb_check["施設一覧"]
types = Counter(ws_check.cell(row=r, column=col_map["施設種別"]).value
                for r in range(4, ws_check.max_row + 1))
print("\n施設種別内訳:")
for k, v in sorted(types.items(), key=lambda x: -x[1]):
    print(f"  {k}: {v}件")
