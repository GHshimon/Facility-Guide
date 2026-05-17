# -*- coding: utf-8 -*-
"""
施設データ.xlsx を新仕様に移行する。
- 「自立1」「自立2」を要介護1の前に挿入（既存データ行は ×）
- 「第3段階」を「第3-1段階」「第3-2段階」に置換（元の第3の値を両列に複写）
行1・行2の結合範囲も列数に合わせて更新する。
"""
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "施設データ.xlsx"


def main():
    wb = load_workbook(XLSX)
    ws = wb["施設一覧"]

    for mc in list(ws.merged_cells.ranges):
        if mc.min_row <= 2:
            try:
                ws.unmerge_cells(str(mc))
            except KeyError:
                pass

    headers = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    if "自立1" in headers:
        print("Already migrated:", XLSX)
        wb.close()
        return

    idx_y1 = headers.index("要介護1") + 1
    ws.insert_cols(idx_y1, 2)
    ws.cell(3, idx_y1).value = "自立1"
    ws.cell(3, idx_y1 + 1).value = "自立2"

    headers2 = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    idx_s4 = headers2.index("第4段階") + 1
    ws.insert_cols(idx_s4, 1)

    headers3 = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    idx_s3 = headers3.index("第3段階") + 1
    ws.cell(3, idx_s3).value = "第3-1段階"
    ws.cell(3, idx_s3 + 1).value = "第3-2段階"

    max_r = ws.max_row
    for r in range(4, max_r + 1):
        ws.cell(r, idx_y1).value = "×"
        ws.cell(r, idx_y1 + 1).value = "×"
        v = ws.cell(r, idx_s3).value
        ws.cell(r, idx_s3 + 1).value = v

    last_c = ws.max_column
    L = get_column_letter(last_c)

    ws.merge_cells(f"A1:{L}1")
    ws["A1"] = "施設データベース（鹿児島県）"
    ws["A1"].font = Font(name="Arial", bold=True, size=14, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws["A1"].fill = PatternFill("solid", start_color="DEEAF1")

    cat = [
        ("A2", "C2", "基本情報", "1F4E79"),
        ("D2", "J2", "受入可能要介護度", "2E75B6"),
        ("K2", "L2", "部屋", "2E75B6"),
        ("M2", "S2", "医療ケア対応", "375623"),
        ("T2", "Y2", "経済条件", "7B3F00"),
        ("Z2", "AE2", "連絡先", "404040"),
        ("AF2", f"{L}2", "相談員生情報", "4B0082"),
    ]
    for start, end, label, color in cat:
        ws.merge_cells(f"{start}:{end}")
        c = ws[start]
        c.value = label
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", start_color=color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb.save(XLSX)
    wb.close()
    print("Migrated:", XLSX, "columns=", last_c)


if __name__ == "__main__":
    main()
